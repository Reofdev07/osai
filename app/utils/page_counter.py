"""Conteo real de páginas (folios) por tipo de archivo. Si no se puede saber, devuelve None: nunca inventa (spec 2026-10-03 §3.9)."""
import os
import re
import zipfile

import fitz


def count_pages(file_path: str, mime_type: str | None) -> int | None:
    mime = (mime_type or "").lower()
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if "pdf" in mime or ext == ".pdf":
            with fitz.open(file_path) as doc:
                return doc.page_count or None
        if mime.startswith("image/"):
            return 1
        if "wordprocessingml" in mime or ext == ".docx":
            return _docx_pages(file_path)
        if "spreadsheetml" in mime or ext == ".xlsx":
            return _xlsx_sheets(file_path)
    except Exception as e:  # archivo dañado o formato inesperado: el funcionario diligencia los folios
        print(f"page_counter: no se pudieron contar las páginas de {os.path.basename(file_path)}: {e}")
    return None


def _docx_pages(path: str) -> int | None:
    with zipfile.ZipFile(path) as z:
        if "docProps/app.xml" not in z.namelist():
            return None
        xml = z.read("docProps/app.xml").decode("utf-8", errors="ignore")
    match = re.search(r"<(?:\w+:)?Pages>(\d+)</(?:\w+:)?Pages>", xml)
    pages = int(match.group(1)) if match else 0
    return pages or None


def _xlsx_sheets(path: str) -> int | None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        return len(workbook.sheetnames) or None
    finally:
        workbook.close()
